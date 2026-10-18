"""sheet-tool: upload an .xlsx file, get it back with one column written by an AI model.

Settings live in one YAML file in S3. Everything else comes from environment variables.
"""
import io
import logging
import os
import re
import time

import boto3
import httpx
import yaml
from botocore.config import Config
from botocore.exceptions import BotoCoreError, ClientError
from fastapi import FastAPI, HTTPException, Request, UploadFile
from fastapi.responses import HTMLResponse, JSONResponse, Response
from openpyxl import load_workbook
from pydantic import BaseModel, Field

# Configuration: only from environment variables, never from files baked into the image.
S3_BUCKET = os.environ["S3_BUCKET"]
S3_KEY = os.getenv("S3_KEY", "settings.yaml")
OPENROUTER_BASE_URL = os.getenv("OPENROUTER_BASE_URL", "https://openrouter.ai/api/v1")
OPENROUTER_API_KEY = os.environ["OPENROUTER_API_KEY"]
MAX_UPLOAD_MB = float(os.getenv("MAX_UPLOAD_MB", "10"))
MAX_UPLOAD_BYTES = int(MAX_UPLOAD_MB * 1024 * 1024)

logging.basicConfig(level=os.getenv("LOG_LEVEL", "INFO"), format="%(levelname)s %(name)s %(message)s")
log = logging.getLogger("sheet-tool")


class SkipHealthChecks(logging.Filter):
    """Kubernetes calls the health endpoints every few seconds; don't log each call."""

    def filter(self, record):
        return "/health/" not in record.getMessage()


logging.getLogger("uvicorn.access").addFilter(SkipHealthChecks())
logging.getLogger("httpx").setLevel(logging.WARNING)  # one line per AI call is too much

# boto3 finds the endpoint, region and credentials itself (AWS_* variables, or the pod's IAM role).
# Short timeouts: a hung S3 call should fail fast, not hold a request for minutes.
s3 = boto3.client("s3", config=Config(connect_timeout=2, read_timeout=5, retries={"max_attempts": 3}))
s3_check = boto3.client("s3", config=Config(connect_timeout=1, read_timeout=2, retries={"max_attempts": 1}))

# httpx also reads HTTPS_PROXY and NO_PROXY, in case the cluster's outbound traffic goes through a proxy.
ai = httpx.Client(
    base_url=OPENROUTER_BASE_URL,
    headers={"Authorization": f"Bearer {OPENROUTER_API_KEY}"},
    timeout=httpx.Timeout(60, connect=5),
)


class Settings(BaseModel):
    model: str = "openai/gpt-4o-mini"
    prompt: str = "Summarise in five words: {text}"
    max_rows: int = Field(500, ge=1, le=5000)


def load_settings(client=s3) -> Settings:
    try:
        body = client.get_object(Bucket=S3_BUCKET, Key=S3_KEY)["Body"].read()
    except ClientError as error:
        if error.response["Error"]["Code"] == "NoSuchKey":
            return Settings()  # nothing saved yet
        raise
    return Settings(**(yaml.safe_load(body) or {}))


def settings_or_503() -> Settings:
    try:
        return load_settings()
    except (BotoCoreError, ClientError) as error:
        log.error("settings storage failed: %r", error)
        raise HTTPException(503, "Settings storage is unreachable")


def save_settings(settings: Settings):
    body = yaml.safe_dump(settings.model_dump(), sort_keys=False).encode()
    s3.put_object(Bucket=S3_BUCKET, Key=S3_KEY, Body=body, ContentType="application/yaml")


def ask_ai(model: str, prompt: str) -> str:
    response = ai.post("/chat/completions", json={"model": model, "messages": [{"role": "user", "content": prompt}]})
    response.raise_for_status()
    return response.json()["choices"][0]["message"]["content"].strip()


app = FastAPI(title="sheet-tool")


@app.middleware("http")
async def limit_upload_size(request: Request, call_next):
    # Refuse a too-large upload before reading it. The ingress should have its own limit too.
    length = request.headers.get("content-length")
    if length and int(length) > MAX_UPLOAD_BYTES:
        return JSONResponse({"detail": f"File is larger than {MAX_UPLOAD_MB:g} MB"}, status_code=413)
    return await call_next(request)


PAGE = """<!doctype html>
<title>sheet-tool</title>
<h1>sheet-tool</h1>
<p>Upload an .xlsx file (up to {limit} MB). Column A is sent to the AI model; the answers come back in a new column.</p>
<form action="process" method="post" enctype="multipart/form-data">
  <input type="file" name="file" accept=".xlsx" required>
  <button>Process and download</button>
</form>
"""


@app.get("/", response_class=HTMLResponse)
def page():
    return PAGE.format(limit=f"{MAX_UPLOAD_MB:g}")


@app.post("/process")
def process(file: UploadFile):
    if not (file.filename or "").lower().endswith(".xlsx"):
        raise HTTPException(415, "Upload an .xlsx file")
    data = file.file.read(MAX_UPLOAD_BYTES + 1)
    if len(data) > MAX_UPLOAD_BYTES:
        raise HTTPException(413, f"File is larger than {MAX_UPLOAD_MB:g} MB")

    settings = settings_or_503()
    try:
        workbook = load_workbook(io.BytesIO(data))
    except Exception:
        raise HTTPException(400, "Not a valid .xlsx file")
    sheet = workbook.active
    rows = sheet.max_row - 1
    if rows > settings.max_rows:
        raise HTTPException(413, f"{rows} rows; the limit is {settings.max_rows}")

    started = time.monotonic()
    answer_column = sheet.max_column + 1
    sheet.cell(row=1, column=answer_column, value="AI answer")
    try:
        for row in range(2, sheet.max_row + 1):
            text = sheet.cell(row=row, column=1).value
            if text is not None:
                prompt = settings.prompt.replace("{text}", str(text))
                sheet.cell(row=row, column=answer_column, value=ask_ai(settings.model, prompt))
    except httpx.HTTPError as error:
        log.error("AI call failed: %r", error)
        raise HTTPException(502, "The AI service failed; try again later")

    output = io.BytesIO()
    workbook.save(output)
    # Log sizes and timings, never file contents or keys.
    log.info("processed rows=%d bytes_in=%d seconds=%.1f", rows, len(data), time.monotonic() - started)

    # The file name comes from the user: keep only safe characters before putting it in a header.
    safe_name = re.sub(r"[^A-Za-z0-9._-]", "_", file.filename)
    return Response(
        output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="processed-{safe_name}"'},
    )


@app.get("/settings")
def get_settings() -> Settings:
    return settings_or_503()


@app.put("/settings")
def put_settings(settings: Settings) -> Settings:
    try:
        save_settings(settings)
    except (BotoCoreError, ClientError) as error:
        log.error("settings storage failed: %r", error)
        raise HTTPException(503, "Settings storage is unreachable")
    log.info("settings saved model=%s max_rows=%d", settings.model, settings.max_rows)
    return settings


@app.get("/health/live")
def live():
    # "The process is running and answering." Checks nothing else, so a broken S3 never causes restarts.
    return {"status": "ok"}


@app.get("/health/ready")
def ready():
    # "I can do my job right now." The app can't work without its settings, so check that S3 answers.
    try:
        load_settings(client=s3_check)
    except Exception as error:
        log.warning("not ready: %r", error)
        raise HTTPException(503, "Settings storage is unreachable")
    return {"status": "ready"}
